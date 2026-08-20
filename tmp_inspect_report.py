from collections import Counter
from pathlib import Path
from docx import Document

source = Path(r"C:\Users\Hp\Downloads\NAISHA SIP REPORT.docx")
doc = Document(source)

print(f"PARAGRAPHS: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip().replace("\t", " ")
    if text:
        print(f"P {i:03d} | {p.style.name} | {text}")

print(f"TABLES: {len(doc.tables)}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: {len(table.rows)}x{len(table.columns)}")
    for row in table.rows:
        print(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))

print("HEADERS")
for si, section in enumerate(doc.sections):
    header = " / ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    footer = " / ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
    if header or footer:
        print(f"SECTION {si + 1} HEADER={header!r} FOOTER={footer!r}")

print("STYLES")
print(Counter(p.style.name for p in doc.paragraphs))
