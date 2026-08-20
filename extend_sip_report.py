from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement


SOURCE = Path(r"C:\Users\Hp\Downloads\NAISHA SIP REPORT.docx")
OUTPUT = Path(r"C:\Users\Hp\Documents\ChatGPT\SG Traders\NAISHA SIP REPORT Extended.docx")


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag.endswith('}pPr'):
            continue
        paragraph._p.remove(child)


def insert_paragraph_before(anchor, text, style_name, template=None, bold=False):
    paragraph = anchor.insert_paragraph_before(style=style_name)
    if template is not None and template._p.pPr is not None:
        existing = paragraph._p.pPr
        if existing is not None:
            paragraph._p.remove(existing)
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    run = paragraph.add_run(text)
    if template is not None and template.runs and template.runs[0]._r.rPr is not None:
        run._r.get_or_add_rPr().append(deepcopy(template.runs[0]._r.rPr))
    if bold:
        run.bold = True
    return paragraph


def set_cell(cell, text):
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.add_run(text)


doc = Document(SOURCE)

# Remove two accidental fragments that belong to a different sentiment-analysis report.
for paragraph in list(doc.paragraphs):
    if paragraph.text.strip().startswith("96.2 percent. Hyperparameter tuning"):
        remove_paragraph(paragraph)
    elif paragraph.text.strip() == "23.7 percent for negative reviews.":
        remove_paragraph(paragraph)

# Refresh the existing manual contents table so it includes the new conclusion material.
contents = doc.tables[0]
for row in contents.rows:
    if len(row.cells) >= 3 and row.cells[0].text.strip() == "5.":
        set_cell(row.cells[2], "21-25")
    if len(row.cells) >= 3 and row.cells[0].text.strip() == "5.3":
        set_cell(row.cells[2], "23-24")
    if len(row.cells) >= 3 and row.cells[0].text.strip() == "6.":
        set_cell(row.cells[2], "26")

bibliography_row_index = next(
    index for index, row in enumerate(contents.rows)
    if row.cells[0].text.strip() == "6."
)
for number, topic, page in [
    ("5.4", "Limitations, Responsible Use and Practical Constraints", "24-25"),
    ("5.5", "Internship Learning and Professional Development", "25"),
]:
    new_row = contents.add_row()
    for cell, value in zip(new_row.cells, [number, topic, page]):
        set_cell(cell, value)
    row_element = new_row._tr
    contents._tbl.remove(row_element)
    contents._tbl.insert(bibliography_row_index, row_element)
    bibliography_row_index += 1

paragraphs = doc.paragraphs
reference_heading = next(
    paragraph for paragraph in paragraphs
    if paragraph.text.strip().upper() == "REFERENCES/BIBLIOGRAPHY"
)
heading_template = next(
    paragraph for paragraph in paragraphs
    if paragraph.text.strip() == "5.3 Future Scope and System Enhancements"
)
body_template = next(
    paragraph for paragraph in paragraphs
    if paragraph.text.strip().startswith("This project shows a complete journey")
)

extension = [
    ("heading", "5.4 Limitations, Responsible Use and Practical Constraints"),
    ("body", "The four projects were designed as focused learning exercises, so their strengths are clarity, simplicity and ease of execution. They are useful for demonstrating distinct AI approaches, but they should not be presented as production-ready systems. A practical AI solution needs a larger test plan, dependable input handling, documentation of assumptions and a process for monitoring behaviour after deployment. Recognising these boundaries was an important part of evaluating the work honestly."),
    ("body", "The rule-based chatbot works well for the set of greetings, questions and exit commands that are explicitly included in the program. Its limitation is that it does not understand the meaning of a sentence beyond those stored conditions. A user may ask a reasonable question in a different form and receive the fallback reply. In a larger application, the next step would be to group similar questions by intent, keep the response rules in a separate configuration file and record unanswered queries for later improvement."),
    ("body", "The breast cancer classifier illustrates the supervised learning workflow, but its result is only an educational experiment on the standard scikit-learn dataset. The program uses one 80:20 split and reports accuracy, which is a helpful starting measure but not a complete clinical evaluation. Before a model could support healthcare decisions, it would require validation on independent data, additional measures such as precision, recall and a confusion matrix, careful review for bias and clear oversight by qualified medical professionals. It must never be treated as a substitute for diagnosis."),
    ("body", "The movie recommendation system uses a small hand-written collection of films and ranks titles by the number of matching genres. This makes the logic easy to inspect, although it does not use ratings, watch history, popularity, language preferences or feedback from other users. The score also treats all matching genres equally. A more mature recommender would use a larger catalogue, collect user feedback with consent and evaluate whether the recommendations genuinely help users discover relevant content."),
    ("body", "The image classifier performs inference with the ImageNet-pretrained MobileNetV2 model. The implementation correctly follows the required 224 x 224 image preparation process, but the model is not fine-tuned on a custom dataset in the repository. Its prediction therefore depends on how closely the supplied image resembles the classes on which MobileNetV2 was trained. The confidence score shown for one image is the model's score for that prediction, not proof of the model's overall accuracy. In addition, the current image path is local to the development machine and would need to be replaced with a file-selection or upload mechanism before sharing the application."),
    ("body", "Responsible use also involves the data that surrounds an AI system. A future version that accepts medical records, conversation logs or user preference data would need informed consent, secure storage, restricted access and a policy for deleting data when it is no longer needed. These requirements are not extras added at the end of development. They influence the choice of data, the interface presented to the user and the way results are explained."),
    ("heading", "5.5 Internship Learning and Professional Development"),
    ("body", "The internship made the difference between four forms of AI much clearer in practice. The chatbot showed how a programmer can define every response directly. The Decision Tree showed how labelled examples can be used to learn decision rules. The recommendation system demonstrated ranking from user preferences, while MobileNetV2 showed how a pre-trained deep learning model can be used once the input is prepared in the correct form. Working through the projects in this order made the increase in complexity feel logical rather than abrupt."),
    ("body", "Another important learning outcome was the value of careful setup and reproducibility. The classification script fixes its random state, the recommendation script normalises genre input before comparison, and the image-classification script applies the preprocessing expected by the model. Small implementation choices like these make a program easier to rerun and easier for another person to understand. The GitHub repository supports this process by keeping each project in a separate folder with a README file and a clearly named Python program."),
    ("body", "Testing was not limited to checking whether a script started successfully. The chatbot needed both recognised and unrecognised inputs. The classifier needed its predicted and actual labels to be compared. The recommendation system needed different genre combinations to confirm that the ranking changed as intended. The image classifier needed a valid image file and an interpretation of its prediction that did not overstate the confidence value. This approach helped connect output checking with the original objective of each project."),
    ("body", "The work also strengthened practical communication skills. Explaining a model to a reader requires more than copying code into a report. It requires stating what the program receives as input, what processing takes place, what the output means and where the result should be used cautiously. Preparing the repository and this report therefore became part of the technical work, not a separate administrative task."),
    ("body", "Overall, the internship provided a strong foundation for future work in Python-based AI development. The next stage is to deepen the evaluation of each system, improve the user interfaces and deploy selected projects in a controlled environment. The portfolio now provides a practical starting point for moving from introductory implementations to better tested applications that are useful, understandable and responsibly designed."),
]

for kind, text in extension:
    if kind == "heading":
        insert_paragraph_before(reference_heading, text, "List Paragraph", heading_template, bold=True)
    else:
        insert_paragraph_before(reference_heading, text, "Body Text", body_template)

# Replace the copied sentiment-analysis bibliography with references relevant to this repository.
references = [
    "[1] N. Jaiswal, DecodeLabs AI Internship Projects, GitHub repository, 2026. Available: https://github.com/naisha11-08/DecodeLabs-AI-Internship",
    "[2] F. Pedregosa et al., Scikit-learn: Machine Learning in Python, Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
    "[3] scikit-learn developers, Breast cancer wisconsin dataset, scikit-learn documentation, 2026. Available: https://scikit-learn.org/stable/datasets/toy_dataset.html#breast-cancer-dataset",
    "[4] L. Breiman, J. H. Friedman, R. A. Olshen and C. J. Stone, Classification and Regression Trees. Belmont, CA: Wadsworth, 1984.",
    "[5] TensorFlow authors, MobileNet models, TensorFlow Keras API documentation, 2026. Available: https://www.tensorflow.org/api_docs/python/tf/keras/applications/MobileNetV2",
    "[6] M. Sandler, A. Howard, M. Zhu, A. Zhmoginov and L. Chen, MobileNetV2: Inverted Residuals and Linear Bottlenecks, Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, pp. 4510-4520, 2018.",
    "[7] Python Software Foundation, Python 3 documentation, 2026. Available: https://docs.python.org/3/",
]
reference_paragraphs = [
    paragraph for paragraph in doc.paragraphs
    if paragraph._p.getprevious() is not None and paragraph.text.strip().startswith(("Bo Pang", "A. Pak", "R. Socher", "M. Hu", "A.Go", "Y. Kim", "J. Devlin"))
]
for paragraph, text in zip(reference_paragraphs, references):
    clear_paragraph(paragraph)
    paragraph.add_run(text)

doc.core_properties.title = "Summer Internship Project Report: Four AI Projects at DecodeLabs"
doc.core_properties.subject = "Extended SIP report based on the DecodeLabs AI Internship repository"
doc.core_properties.comments = ""
doc.save(OUTPUT)
print(OUTPUT)
